"""Import structured test prompts from Excel (.xlsx) — Dassault / QA template compatible."""

from __future__ import annotations

import io
import re
from typing import Any

from src.common.models import StructuredTestPrompt, TestEnvironment

# Header aliases (case-insensitive, normalized to snake key)
_HEADER_ALIASES: dict[str, str] = {
    "test_id": "test_id",
    "test id": "test_id",
    "id": "test_id",
    "case_id": "test_id",
    "site_url": "site_url",
    "site url": "site_url",
    "url": "site_url",
    "application_url": "site_url",
    "feature": "feature",
    "module": "feature",
    "feature_module": "feature",
    "test_name": "test_name",
    "test name": "test_name",
    "name": "test_name",
    "title": "test_name",
    "objective": "objective",
    "test_objective": "objective",
    "expected_outcome": "expected_outcome",
    "expected outcome": "expected_outcome",
    "expected result": "expected_outcome",
    "expected": "expected_outcome",
    "environment": "environment",
    "env": "environment",
    "owner_team": "owner_team",
    "owner team": "owner_team",
    "team": "owner_team",
    "steps": "steps",
    "test steps": "steps",
    "step": "steps",
    "procedure": "steps",
}


def _norm_header(cell: Any) -> str:
    if cell is None:
        return ""
    return re.sub(r"\s+", " ", str(cell).strip().lower())


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_steps(raw: str, row: dict[str, Any], step_col_prefix: str = "step") -> list[str]:
    lines: list[str] = []
    if raw:
        for line in re.split(r"[\r\n]+", raw):
            line = line.strip()
            if not line:
                continue
            line = re.sub(r"^\d+[\).\s-]+", "", line)
            if line:
                lines.append(line)
    if lines:
        return lines
    # Step 1, Step 2, ... columns (keys like step_1 or step1)
    for key in sorted(row.keys()):
        if not key.startswith(step_col_prefix):
            continue
        if row[key]:
            text = _cell_str(row[key])
            if text:
                lines.append(re.sub(r"^\d+[\).\s-]+", "", text))
    return lines


def _map_environment(raw: str) -> TestEnvironment:
    low = raw.lower()
    if "prod" in low:
        return TestEnvironment.PROD
    if "stage" in low or "staging" in low:
        return TestEnvironment.STAGE
    return TestEnvironment.DEVELOP


def parse_excel_bytes(data: bytes, *, sheet_index: int = 0) -> list[StructuredTestPrompt]:
    """
    Parse .xlsx into StructuredTestPrompt list.

    Expects a header row. Required columns: test_id, site_url, feature, test_name,
    objective, expected_outcome, and steps (multiline cell or Step 1..N columns).
    """
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = wb.worksheets
    if sheet_index >= len(sheets):
        raise ValueError(f"Sheet index {sheet_index} out of range")
    ws = sheets[sheet_index]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise ValueError("Excel sheet is empty") from exc

    col_map: dict[int, str] = {}
    for idx, cell in enumerate(header_row):
        norm = _norm_header(cell)
        alias = _HEADER_ALIASES.get(norm)
        if alias:
            col_map[idx] = alias
        elif norm.startswith("step"):
            col_map[idx] = norm.replace(" ", "_")

    required = {"test_id", "site_url", "feature", "test_name", "objective", "expected_outcome"}
    found = set(col_map.values())
    has_step_columns = "steps" in found or any(k.startswith("step") for k in found)
    missing = required - found
    if not has_step_columns:
        missing.add("steps")
    else:
        missing.discard("steps")
    if missing:
        raise ValueError(
            f"Missing required column(s): {', '.join(sorted(missing))}. "
            f"Found headers mapped to: {sorted(found)}"
        )

    prompts: list[StructuredTestPrompt] = []
    for row_cells in rows_iter:
        if not any(row_cells):
            continue
        row_dict: dict[str, Any] = {}
        for idx, val in enumerate(row_cells):
            if idx in col_map:
                row_dict[col_map[idx]] = val

        test_id = _cell_str(row_dict.get("test_id"))
        if not test_id:
            continue

        steps_raw = _cell_str(row_dict.get("steps", ""))
        steps = _parse_steps(steps_raw, row_dict)
        if not steps:
            raise ValueError(f"Row test_id={test_id}: no steps found")

        env_raw = _cell_str(row_dict.get("environment", "develop"))
        owner = _cell_str(row_dict.get("owner_team", "")) or None

        prompts.append(
            StructuredTestPrompt(
                test_id=test_id,
                site_url=_cell_str(row_dict.get("site_url")),
                feature=_cell_str(row_dict.get("feature")),
                test_name=_cell_str(row_dict.get("test_name")),
                objective=_cell_str(row_dict.get("objective")),
                expected_outcome=_cell_str(row_dict.get("expected_outcome")),
                environment=_map_environment(env_raw) if env_raw else TestEnvironment.DEVELOP,
                owner_team=owner,
                steps=steps,
            )
        )

    wb.close()
    if not prompts:
        raise ValueError("No data rows found in Excel file")
    return prompts
