"""Generate tests/fixtures/sample_import.xlsx for Excel import UI/API testing."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

OUT = Path(__file__).resolve().parent / "sample_import.xlsx"

HEADERS = [
    "Test ID",
    "Site URL",
    "Feature",
    "Test Name",
    "Objective",
    "Expected Outcome",
    "Environment",
    "Owner Team",
    "Steps",
]

ROWS = [
    [
        "TC01_login_success",
        "https://www.saucedemo.com/",
        "login",
        "Valid user login",
        "Verify standard user can log in and reach inventory",
        "Products page visible after login",
        "develop",
        "auth-frontend",
        "Fill username with standard_user\n"
        "Fill password with secret_sauce\n"
        "Click Login\n"
        "Verify URL contains inventory.html\n"
        "Verify text Products is visible",
    ],
    [
        "TC10_intentional_fail",
        "https://www.saucedemo.com/",
        "login",
        "Intentional failure demo",
        "Demonstrate fail report and notify routing",
        "This step should not pass",
        "develop",
        "auth-frontend",
        "Fill username with standard_user\n"
        "Fill password with wrong_password\n"
        "Click Login\n"
        "Verify text This text does not exist on page is visible",
    ],
]


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "TestCases"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    for row in ROWS:
        ws.append(row)

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 24
    ws.column_dimensions["E"].width = 36
    ws.column_dimensions["F"].width = 32
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 16
    ws.column_dimensions["I"].width = 48

    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
